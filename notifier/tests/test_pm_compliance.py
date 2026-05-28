"""Tests for src/pm_compliance.py — pure, no network required."""

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.pm_compliance import (
    PMSummary,
    build_pm_summary_for_machine,
    load_pm_compliance,
)

FIXTURE = Path(__file__).parent / "fixtures" / "sample_pm_compliance.xlsx"


@pytest.fixture()
def pm_df():
    return load_pm_compliance(FIXTURE.read_bytes())


# ── load_pm_compliance ────────────────────────────────────────────────────────

def test_load_basic(pm_df):
    assert "WC Object ID" in pm_df.columns
    assert "Core PM" in pm_df.columns
    assert "Urgency" in pm_df.columns
    assert len(pm_df) == 7


def test_load_missing_column_raises():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["WC Object ID", "Core PM"])   # missing Urgency
    ws2.append(["10009999", "Lubrication"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Urgency"):
        load_pm_compliance(buf.getvalue())


def test_load_urgency_coercion():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["WC Object ID", "Core PM", "Urgency"])
    ws2.append(["10009999", "Lubrication", "N/A"])   # non-numeric → 0
    ws2.append(["10009999", "Inspection",  1])
    buf = BytesIO()
    wb.save(buf)
    df = load_pm_compliance(buf.getvalue())
    assert list(df["Urgency"]) == [0, 1]


def test_load_strips_whitespace():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["WC Object ID", "Core PM", "Urgency"])
    ws2.append([" 10009999 ", " Lubrication ", 1])
    buf = BytesIO()
    wb.save(buf)
    df = load_pm_compliance(buf.getvalue())
    assert df["WC Object ID"].iloc[0] == "10009999"
    assert df["Core PM"].iloc[0] == "Lubrication"


# ── build_pm_summary_for_machine ──────────────────────────────────────────────

def test_build_returns_none_no_urgency_rows(pm_df):
    # 10009999 Inspection/Check Guards has Urgency=0; if it were the only row → None
    # Use a WC that has only Urgency=0 rows — craft minimal df
    import pandas as pd
    df = pd.DataFrame({
        "WC Object ID": ["ABC123"],
        "Core PM": ["Lubrication"],
        "Urgency": [0],
    })
    assert build_pm_summary_for_machine(df, "ABC123") is None


def test_build_returns_none_unknown_machine(pm_df):
    assert build_pm_summary_for_machine(pm_df, "99999999") is None


def test_build_counts_correctly(pm_df):
    result = build_pm_summary_for_machine(pm_df, "10009999")
    assert result is not None
    assert result.total_overdue == 4   # Lubrication×3 + Inspection×1 (Urgency=1 only)


def test_build_sorted_desc(pm_df):
    result = build_pm_summary_for_machine(pm_df, "10009999")
    assert result is not None
    counts = [r.overdue_count for r in result.by_core_pm]
    assert counts == sorted(counts, reverse=True)


def test_build_multiple_core_pms(pm_df):
    result = build_pm_summary_for_machine(pm_df, "10009999")
    assert result is not None
    core_pm_names = {r.core_pm for r in result.by_core_pm}
    assert "Lubrication" in core_pm_names
    assert "Inspection" in core_pm_names
    lube = next(r for r in result.by_core_pm if r.core_pm == "Lubrication")
    assert lube.overdue_count == 3
    insp = next(r for r in result.by_core_pm if r.core_pm == "Inspection")
    assert insp.overdue_count == 1


def test_build_single_machine_cleaning(pm_df):
    result = build_pm_summary_for_machine(pm_df, "10009998")
    assert result is not None
    assert result.total_overdue == 1
    assert result.by_core_pm[0].core_pm == "Cleaning"


def test_build_calibration_machine(pm_df):
    result = build_pm_summary_for_machine(pm_df, "20005555")
    assert result is not None
    assert result.total_overdue == 1
    assert result.by_core_pm[0].core_pm == "Calibration"
