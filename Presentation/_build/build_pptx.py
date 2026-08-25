# -*- coding: utf-8 -*-
"""Builds Presentation/KPI_BSP_Working_Session.pptx from content.py."""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(__file__))
import content as C

# Company brand palette: Green #006548, Light Green #76BC21, Black #3D3935
NAVY = RGBColor(0x00, 0x65, 0x48)      # brand Green — primary (headers, dark fills)
ACCENT = RGBColor(0x76, 0xBC, 0x21)    # brand Light Green — accent (kickers, highlights)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x3D, 0x39, 0x35)     # brand Black — primary body text
LIGHT_BG = RGBColor(0xEA, 0xF2, 0xE2)  # pale tint of brand Green, for table zebra rows
CALLOUT_BG = RGBColor(0xF0, 0xF7, 0xE3)  # pale tint of brand Light Green, for callout boxes
HIGHLIGHT_BG = RGBColor(0xD9, 0xEC, 0xB2)  # deeper Light Green tint, for in-table emphasis
GRAY = RGBColor(0x74, 0x70, 0x6C)      # muted brand Black, for footers/italic notes
GREEN = ACCENT
RED = RGBColor(0xB0, 0x2A, 0x2A)       # reserved for negative/warning callouts (non-brand)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OUT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "KPI_BSP_Working_Session.pptx",
)


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def add_bg(slide, color=WHITE):
    fill_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    fill_rect.fill.solid()
    fill_rect.fill.fore_color.rgb = color
    fill_rect.line.fill.background()
    fill_rect.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(fill_rect._element)
    spTree.insert(2, fill_rect._element)
    return fill_rect


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                 color=NAVY, align=PP_ALIGN.LEFT, italic=False, font="Calibri",
                 anchor=None, line_spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
            run.font.name = font
    return box


def add_bullets(slide, left, top, width, height, items, size=14, color=BLACK,
                 bullet_char="•", space_after=6, bold_lead=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            lead, rest = item
            p.text = ""
            r1 = p.add_run()
            r1.text = f"{bullet_char} {lead}: "
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.color.rgb = NAVY
            r2 = p.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
        else:
            p.text = f"{bullet_char} {item}"
            for run in p.runs:
                run.font.size = Pt(size)
                run.font.color.rgb = color
        p.space_after = Pt(space_after)
    return box


def add_footer(slide, text, page_no=None):
    add_textbox(slide, Inches(0.5), Inches(7.15), Inches(9), Inches(0.3), text,
                size=9, color=GRAY, italic=True)
    if page_no:
        add_textbox(slide, Inches(12.4), Inches(7.15), Inches(0.6), Inches(0.3),
                    str(page_no), size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def add_section_header(slide, kicker, title):
    add_textbox(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.35), kicker,
                size=13, bold=True, color=ACCENT)
    add_textbox(slide, Inches(0.5), Inches(0.62), Inches(12.3), Inches(0.75), title,
                size=27, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.35), Inches(12.3), Pt(2.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    line.shadow.inherit = False


def add_box(slide, left, top, width, height, fill=LIGHT_BG, line_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    return shp


def add_flow_chevrons(slide, left, top, width, height, labels, size=12.5):
    n = len(labels)
    gap = Pt(6)
    each_w = int((width - gap * (n - 1)) / n)
    x = left
    for i, label in enumerate(labels):
        shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, top, Emu(each_w), height)
        shp.fill.solid()
        shp.fill.fore_color.rgb = NAVY if i % 2 == 0 else ACCENT
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = WHITE
        x = Emu(int(x) + each_w - Emu(Inches(0.12)))


def kpi_table(slide, left, top, width, rows, col_widths=None, header=True,
              font_size=10.5, header_size=10.5, row_h=Inches(0.34)):
    nrows = len(rows)
    ncols = len(rows[0])
    height = row_h * nrows
    gtable = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            gtable.columns[i].width = w
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            cell = gtable.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Pt(4)
            cell.margin_right = Pt(4)
            cell.margin_top = Pt(2)
            cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            for run in para.runs:
                run.font.size = Pt(header_size if (header and r == 0) else font_size)
                run.font.bold = header and r == 0
                run.font.color.rgb = WHITE if (header and r == 0) else BLACK
            if header and r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT_BG
    return gtable


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ---- Slide 1: Title ----
    s = blank_slide(prs)
    add_bg(s, NAVY)
    add_textbox(s, Inches(1), Inches(2.5), Inches(11.3), Inches(1.4), C.PPT_TITLE,
                size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s, Inches(1.3), Inches(3.75), Inches(10.7), Inches(1.0), C.PPT_SUBTITLE,
                size=18, color=RGBColor(0xCF, 0xE7, 0xB8), align=PP_ALIGN.CENTER, italic=True)
    add_textbox(s, Inches(1.3), Inches(6.6), Inches(10.7), Inches(0.5),
                "Folding Carton Actionable Intelligence Platform",
                size=13, color=RGBColor(0x9E, 0xC9, 0x7A), align=PP_ALIGN.CENTER)

    # ---- Slide 2: Purpose ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "WORKING SESSION", "Purpose of Today's Session")
    add_bullets(s, Inches(0.6), Inches(1.7), Inches(11.8), Inches(2.2), [
        "Review exactly how each KPI and its BSP (Best Shown Performance) benchmark "
        "are calculated today — in plain business language, not code.",
        "Every formula shown was taken directly from the production calculation "
        "script — nothing here is invented.",
        "This is a working session, not a final declaration that the current "
        "methodology is correct.",
    ], size=17, space_after=14)
    add_box(s, Inches(0.6), Inches(4.2), Inches(11.9), Inches(2.3), fill=CALLOUT_BG)
    add_textbox(s, Inches(0.9), Inches(4.4), Inches(11.3), Inches(0.4),
                "Pilot Scope — keep in mind throughout", size=15, bold=True, color=NAVY)
    add_bullets(s, Inches(0.9), Inches(4.85), Inches(11.3), Inches(1.6), [
        "Only 13 machines are currently in scope; all others are excluded.",
        "History is available from January 1, 2025 only — benchmarks are designed "
        "to use up to 24 months, but currently run on less than that.",
    ], size=13.5, space_after=6)

    # ---- Slide 3: Overall Calculation Flow ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "HOW IT WORKS", "Overall Calculation Flow")
    add_flow_chevrons(s, Inches(0.4), Inches(2.6), Inches(12.6), Inches(1.0), C.BSP_FLOW_STEPS, size=12)
    add_bullets(s, Inches(0.7), Inches(4.1), Inches(11.8), Inches(2.6), [
        ("Actual", "what the machine did last week."),
        ("BSP", "the benchmark, built from that machine's own best demonstrated history."),
        ("BSP Performance / Gap", "the difference between Actual and BSP."),
        ("Sheet Gap", "that gap converted into an estimated number of sheets lost — the common scale used to rank every opportunity."),
    ], size=15, space_after=10)

    # ---- Slide 4: BSP in one picture ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "THE CORE MECHANISM", "How BSP Is Calculated")
    steps_short = [
        "History pool: every qualifying run, ~last 24 months, run time > 2 hrs",
        "Percentile, not best-ever: P75 (higher-better KPIs) or P25 (lower-better KPIs)",
        "3 levels of specificity: Die (≥15 runs) → Board (≥25) → Machine (≥35)",
        "Most specific level with enough data wins — sets HIGH / MEDIUM / LOW confidence",
        "Die-level BSPs roll up to one machine-week BSP, weighted by scheduled hours",
        "Coverage gate: insight only fires if >50% of the week's hours have a resolved BSP",
    ]
    y = Inches(1.65)
    for i, txt in enumerate(steps_short, start=1):
        add_box(s, Inches(0.6), y, Inches(0.55), Inches(0.75), fill=NAVY)
        add_textbox(s, Inches(0.6), y, Inches(0.55), Inches(0.75), str(i), size=20, bold=True,
                    color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_textbox(s, Inches(1.3), y, Inches(11.2), Inches(0.75), txt, size=14.5,
                    color=BLACK, anchor=MSO_ANCHOR.MIDDLE)
        y = Emu(int(y) + int(Inches(0.83)))

    # ---- Slide 4b: BSP walked example — finding the percentile (Speed) ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "WORKED EXAMPLE", "BSP Step-by-Step — Speed: Finding the Percentile")
    wt = C.SPEED_BSP_WALKTHROUGH
    add_textbox(s, Inches(0.6), Inches(1.5), Inches(6.4), Inches(1.1),
                wt["pool_intro"], size=13, color=BLACK)
    add_textbox(s, Inches(0.6), Inches(2.65), Inches(6.4), Inches(1.9),
                wt["percentile_note"], size=13, color=BLACK)
    add_box(s, Inches(0.6), Inches(4.65), Inches(6.4), Inches(1.85), fill=CALLOUT_BG)
    add_textbox(s, Inches(0.85), Inches(4.8), Inches(5.9), Inches(0.4),
                "Result", size=14, bold=True, color=NAVY)
    add_textbox(s, Inches(0.85), Inches(5.25), Inches(5.9), Inches(1.1),
                f"BSP Speed (Level 1) = {wt['l1_bsp']:,} sheets/hr\n"
                "This is the speed the machine has already sustained across its "
                "own best quarter of runs — not its single fastest run.",
                size=13, color=BLACK)

    # Sorted-pool table on the right, split into two columns for readability
    speeds = wt["sorted_speeds"]
    rows_tbl = [["Rank", "Speed (sheets/hr)"]]
    for i, v in enumerate(speeds, start=1):
        marker = "  ← P75 zone" if i in (11, 12) else ""
        rows_tbl.append([str(i), f"{v:,}{marker}"])
    gtable = kpi_table(s, Inches(7.5), Inches(1.5), Inches(5.25), rows_tbl,
                        col_widths=[Inches(1.1), Inches(4.15)], font_size=11,
                        header_size=11.5, row_h=Inches(0.335))
    # Highlight the two rows that bracket the 75th percentile
    for r in (11, 12):
        for c in range(2):
            cell = gtable.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HIGHLIGHT_BG

    # ---- Slide 4c: BSP walked example — level fallback + Sheet Gap ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "WORKED EXAMPLE", "BSP Step-by-Step — Speed: Level Fallback & Sheet Gap")
    add_textbox(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.4),
                "This week, Level 1 had exactly enough runs and resolved directly:",
                size=14, bold=True, color=NAVY)
    lvl_rows = [["Level", "Grain", "Qualifying Runs", "Minimum Needed", "Resolved?"]]
    for lvl, grain, runs, minimum, resolved, *_ in wt["levels"]:
        lvl_rows.append([lvl, grain, str(runs), str(minimum), resolved])
    kpi_table(s, Inches(0.6), Inches(2.0), Inches(12.1), lvl_rows,
              col_widths=[Inches(2.5), Inches(4.1), Inches(2.0), Inches(1.8), Inches(1.7)],
              font_size=11.5, header_size=12, row_h=Inches(0.42))
    add_textbox(s, Inches(0.6), Inches(3.9), Inches(12.1), Inches(0.6),
                wt["resolved_summary"], size=12.5, italic=True, color=GRAY)

    fb = wt["fallback_illustration"]
    add_box(s, Inches(0.6), Inches(4.55), Inches(12.1), Inches(1.55), fill=LIGHT_BG)
    add_textbox(s, Inches(0.85), Inches(4.65), Inches(11.6), Inches(0.35),
                fb["title"], size=13, bold=True, color=NAVY)
    fb_line = "   |   ".join(
        f"{lvl}: {runs} runs vs. {minimum} needed → {status}"
        for lvl, runs, minimum, status in fb["rows"]
    )
    add_textbox(s, Inches(0.85), Inches(5.0), Inches(11.6), Inches(0.5), fb_line,
                size=11, color=BLACK)
    add_textbox(s, Inches(0.85), Inches(5.5), Inches(11.6), Inches(0.55), fb["note"],
                size=10.5, italic=True, color=GRAY)

    add_box(s, Inches(0.6), Inches(6.25), Inches(12.1), Inches(0.95), fill=CALLOUT_BG)
    add_textbox(s, Inches(0.85), Inches(6.35), Inches(11.6), Inches(0.75),
                f"Actual Speed = {wt['actual_speed']:,} sheets/hr, Run Hours = {wt['run_hours']}   →   "
                f"Sheet Gap = {wt['gap_calc']}",
                size=13.5, bold=True, color=NAVY)

    # ---- Slide 5: KPI Overview table ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "WHAT WE MEASURE", "KPI Overview — 10 Published KPIs")
    rows = [["KPI", "Category", "Direction"]]
    for kpi in C.KPIS:
        rows.append([kpi["name"], kpi["category"].replace("Lever - Downtime % (sub-reason of Downtime %)", "Sub-lever of Downtime %"), kpi["direction"].replace(" is better", "")])
    kpi_table(s, Inches(0.6), Inches(1.6), Inches(12.1), rows,
              col_widths=[Inches(4.6), Inches(4.5), Inches(3.0)], font_size=11, header_size=12, row_h=Inches(0.47))

    # ---- KPI detail slides (grouped) ----
    def kpi_detail_slide(kicker, title, kpi):
        s = blank_slide(prs)
        add_bg(s)
        add_section_header(s, kicker, title)
        left_w = Inches(7.6)
        add_textbox(s, Inches(0.6), Inches(1.55), left_w, Inches(0.35),
                    f"Direction: {kpi['direction']}", size=13, bold=True, color=ACCENT)
        add_textbox(s, Inches(0.6), Inches(1.95), left_w, Inches(0.7),
                    kpi["measures"], size=13, color=BLACK)
        add_bullets(s, Inches(0.6), Inches(2.75), left_w, Inches(0.6),
                    [("Formula", kpi["formula"])], size=13)
        add_bullets(s, Inches(0.6), Inches(3.35), left_w, Inches(0.6),
                    [("BSP", kpi["bsp"] if len(kpi["bsp"]) < 220 else kpi["bsp"][:210] + "...")], size=12.5)
        add_bullets(s, Inches(0.6), Inches(4.35), left_w, Inches(0.6),
                    [("Sheet Gap", kpi["sheet_gap_formula"])], size=12.5)
        filt = kpi["filters"]
        add_bullets(s, Inches(0.6), Inches(5.15), left_w, Inches(1.9),
                    [f[:150] for f in filt], size=11.5, bullet_char="—")

        # Right side: example box
        ex = kpi["example"]
        add_box(s, Inches(8.5), Inches(1.55), Inches(4.25), Inches(5.5), fill=CALLOUT_BG)
        add_textbox(s, Inches(8.75), Inches(1.7), Inches(3.8), Inches(0.4),
                    "Worked Example", size=15, bold=True, color=NAVY)
        ey = Inches(2.25)
        for key in ("actual", "bsp", "inputs", "calc", "result"):
            if key in ex:
                add_textbox(s, Inches(8.75), ey, Inches(3.8), Inches(0.85), ex[key],
                            size=11.5, color=BLACK)
                ey = Emu(int(ey) + int(Inches(0.75)))
        return s

    kpi_detail_slide("OUTCOME", "OEE — The Headline Number", C.KPIS[0])
    kpi_detail_slide("LEVER", "Speed", C.KPIS[1])
    kpi_detail_slide("LEVER", "Downtime %", C.KPIS[2])
    kpi_detail_slide("SUB-LEVER", "Downtime Reason Breakdown", C.KPIS[9])
    kpi_detail_slide("LEVER", "Avg MR Time (Changeover Time)", C.KPIS[4])

    # Blanket Wash / Feeder Trip combined slide
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "LEVERS", "Blanket Wash & Feeder Trip — Duration and Frequency")
    add_bullets(s, Inches(0.6), Inches(1.6), Inches(12), Inches(1.2), [
        "Two events, each measured two ways: how long it takes (Avg Time) and how "
        "often it happens (per 10K sheets produced).",
        "Identified from an approved, exact list of reason-code descriptions — "
        "not a keyword search — to avoid false matches across plants.",
        "Excluded from the general Downtime Reason KPI so time is never counted twice.",
    ], size=14.5, space_after=10)
    rows = [
        ["KPI", "Formula", "BSP Pool"],
        ["Avg Blanket Wash Time", "Wash Hours ÷ Wash Events", "Runs with ≥1 wash event"],
        ["Avg Feeder Trip Time", "Trip Hours ÷ Trip Events", "Runs with ≥1 trip event"],
        ["Blanket Washes per 10K", "(Wash Events ÷ Total Sheets) × 10,000", "All qualifying runs, incl. 0-event runs"],
        ["Feeder Trips per 10K", "(Trip Events ÷ Total Sheets) × 10,000", "All qualifying runs, incl. 0-event runs"],
    ]
    kpi_table(s, Inches(0.6), Inches(3.1), Inches(12.1), rows,
              col_widths=[Inches(3.3), Inches(5.0), Inches(3.8)], font_size=12, header_size=12.5, row_h=Inches(0.55))

    # ---- Scrap special logic slide ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "SPECIAL BSP LOGIC", "Scrap Loss → Expected Scrap → BSP Performance")
    add_textbox(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.7),
                "Scrap does not use one flat target per machine — it uses an "
                "order-size-aware \"expected scrap\" benchmark instead.",
                size=15, bold=True, color=NAVY)
    add_flow_chevrons(s, Inches(0.4), Inches(2.35), Inches(12.6), Inches(0.75),
                       ["Bucket order by size", "P25 scrap % per bucket (6 mo.)", "Expected = Bucket % × Order Sheets", "Compare to Actual sheets scrapped"], size=11.5)
    add_box(s, Inches(0.6), Inches(3.4), Inches(12.1), Inches(3.5), fill=CALLOUT_BG)
    add_textbox(s, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.4),
                "Worked Example", size=15, bold=True, color=NAVY)
    ex = C.SCRAP_SPECIAL_EXAMPLE
    ey = Inches(4.05)
    for key in ("setup", "bench", "expected", "actual", "gap"):
        add_textbox(s, Inches(0.9), ey, Inches(11.5), Inches(0.5), ex[key], size=14,
                    color=BLACK)
        ey = Emu(int(ey) + int(Inches(0.53)))

    # ---- Sheet Gap slide ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "COMMON CURRENCY", "Sheet Gap — Turning Every KPI Into Sheets")
    add_textbox(s, Inches(0.6), Inches(1.55), Inches(7.7), Inches(1.0),
                C.SHEET_GAP_INTRO, size=13.5, color=BLACK)
    add_bullets(s, Inches(0.6), Inches(2.75), Inches(7.7), Inches(3.6), [
        ("No double-counting", "OEE's own Sheet Gap is the machine's Total Sheet Impact; lever rows explain WHY, they are not added on top."),
        ("Gated", "Only counted when the gap is positive and BSP coverage exceeds 50%."),
        ("Ranking", "Opportunities within a plant are ranked by Sheet Impact, largest first."),
        ("Streak is display-only", "Consecutive-week or 4-week-above-BSP streaks are shown but do not change the Sheet Gap number."),
    ], size=13, space_after=10)
    add_box(s, Inches(8.6), Inches(1.55), Inches(4.15), Inches(4.8), fill=CALLOUT_BG)
    add_textbox(s, Inches(8.85), Inches(1.7), Inches(3.7), Inches(0.4),
                "Worked Example (OEE)", size=15, bold=True, color=NAVY)
    sg = C.SHEET_GAP_EXAMPLE
    ey = Inches(2.25)
    for key in ("line1", "line2", "calc", "line3"):
        add_textbox(s, Inches(8.85), ey, Inches(3.7), Inches(1.0), sg[key], size=12.5,
                    color=BLACK)
        ey = Emu(int(ey) + int(Inches(0.95)))

    # ---- Filters/Exclusions slide ----
    s = blank_slide(prs)
    add_bg(s)
    add_section_header(s, "GUARDRAILS", "Key Filters & Exclusions")
    add_bullets(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.3), [
        "Only scheduled/crewed time counts toward time-based KPIs; runs of 2 hours or less are excluded from every benchmark pool.",
        "Setup Down time counts as machine downtime (Downtime %), but is excluded from reason-level downtime — reason codes attribute only Down time.",
        "MRO/changeover reason codes and the approved Blanket Wash / Feeder Trip reason list are both excluded from the general Downtime Reason KPI, so nothing is counted twice.",
        "Scrap Loss (actual and benchmark) is restricted to single-setup orders only.",
        "Every insight requires BSP coverage strictly greater than 50% of the week's scheduled (or, for Scrap, sheet) volume before it is published.",
    ], size=15.5, space_after=16)

    # ---- Discussion slide ----
    s = blank_slide(prs)
    add_bg(s, NAVY)
    add_textbox(s, Inches(0.6), Inches(0.5), Inches(12), Inches(0.8),
                "Plant Discussion", size=32, bold=True, color=WHITE)
    add_bullets(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.6), [
        "Are these the right KPI definitions — are we measuring what we intend to measure?",
        "Are the right events included, and are the right events excluded?",
        "Are the current filters correct for your plant and machines?",
        "Is the BSP target (percentile + minimum pool size) reasonable and appropriately ambitious?",
        "Is the BSP performance / Sheet Gap comparison method fair?",
        "Should any KPI use different special logic (e.g. expected scrap)?",
        "Does Sheet Gap represent plant performance the way it should?",
        "See the full \"Items for Plant Review\" list in the reference document for 13 specific discussion points.",
    ], size=16, color=RGBColor(0xDD, 0xEE, 0xC7), space_after=12)

    prs.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    build()
